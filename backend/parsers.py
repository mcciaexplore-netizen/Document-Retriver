"""Format-specific extraction. Every record retains a concrete source coordinate."""
import csv
import io
import re
import zipfile
from datetime import date, datetime
from pathlib import Path
import fitz
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pptx import Presentation

MAX_RECORDS = 100_000
MAX_ROWS = 50_000


def numeric(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = re.sub(r"[₹$€£,\s]", "", str(value))
    if re.fullmatch(r"\(?-?\d+(?:\.\d+)?\)?", text):
        return float(text.strip("()")) * (-1 if text.startswith("(") else 1)
    return None


def display(value, number_format=""):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (int, float)):
        if "₹" in number_format:
            number = f"{abs(value):.0f}"
            tail, head = number[-3:], number[:-3]
            groups = []
            while head:
                groups.insert(0, head[-2:])
                head = head[:-2]
            return ("-" if value < 0 else "") + "₹" + ",".join(groups + [tail])
        if "%" in number_format:
            return f"{value * 100:g}%"
        return str(int(value)) if value == int(value) else str(value)
    return str(value)


def validate_container(path, file_type):
    if file_type in {"xlsx", "pptx"}:
        try:
            with zipfile.ZipFile(path) as package:
                members = package.infolist()
                if len(members) > 8000 or sum(m.file_size for m in members) > 200 * 1024 * 1024:
                    raise ValueError("The expanded document exceeds the 200 MB processing limit.")
                if any(m.flag_bits & 1 for m in members):
                    raise ValueError("Password-protected documents are not supported.")
                names = {m.filename for m in members}
                marker = "xl/workbook.xml" if file_type == "xlsx" else "ppt/presentation.xml"
                if marker not in names:
                    raise ValueError("The file contents do not match its extension.")
        except zipfile.BadZipFile as exc:
            raise ValueError("This is not a valid Office document.") from exc
    if file_type == "pdf":
        with open(path, "rb") as handle:
            if not handle.read(1024).lstrip().startswith(b"%PDF-"):
                raise ValueError("This is not a valid PDF document.")


def cell_record(kind, value, header, row, column, coordinate, sheet, metadata, row_context, raw=None):
    return dict(record_type=kind, content=f"{header}: {value}", value=value, header=header[:500], numeric_value=numeric(raw if raw is not None else value), sheet_name=sheet, row_number=row, column_name=column, cell_coordinate=coordinate, metadata_json=metadata, context=row_context)


def parse_xlsx(path):
    formulas = load_workbook(path, data_only=False, keep_links=False)
    values = load_workbook(path, data_only=True, keep_links=False)
    records = []
    try:
        for sheet in formulas:
            sheet_record_start = len(records)
            if sheet.max_row > MAX_ROWS or sheet.max_column > 500:
                raise ValueError("A sheet exceeds the 50,000 row or 500 column processing limit.")
            value_sheet = values[sheet.title]
            header_row = next((row for row in sheet.iter_rows() if any(c.value is not None for c in row)), None)
            if header_row is None:
                continue
            headers = {cell.column: display(cell.value) or get_column_letter(cell.column) for cell in header_row}
            for row in sheet.iter_rows(min_row=header_row[0].row + 1):
                if all(c.value is None for c in row):
                    continue
                rendered = {cell.column: display(value_sheet.cell(cell.row, cell.column).value if cell.data_type == "f" and value_sheet.cell(cell.row, cell.column).value is not None else cell.value, cell.number_format) for cell in row}
                context = " | ".join(f"{headers.get(col, get_column_letter(col))}: {value}" for col, value in rendered.items() if value)
                for cell in row:
                    if cell.value is None:
                        continue
                    formula = str(cell.value) if cell.data_type == "f" else None
                    merged = next((str(rng) for rng in sheet.merged_cells.ranges if cell.coordinate in rng), None)
                    raw = value_sheet.cell(cell.row, cell.column).value if formula else cell.value
                    metadata = {"formula": formula, "number_format": cell.number_format, "merged_range": merged, "header_row": header_row[0].row, "cached_value_available": bool(formula and raw is not None)}
                    records.append(cell_record("spreadsheet_cell", rendered[cell.column], headers.get(cell.column, cell.column_letter), cell.row, cell.column_letter, cell.coordinate, sheet.title, metadata, context, raw))
                    if len(records) > MAX_RECORDS:
                        raise ValueError("The workbook exceeds the 100,000 searchable record limit.")
            if len(records) > sheet_record_start:
                records[sheet_record_start]["metadata_json"]["table_headers"] = [{"key": get_column_letter(column), "label": header} for column, header in headers.items()]
        return records, [], []
    finally:
        formulas.close()
        values.close()


def parse_csv(path):
    raw = Path(path).read_bytes()
    if b"\x00" in raw[:8192] and not raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        raise ValueError("CSV must contain text, not binary data.")
    encoding = "utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
    try:
        text = raw.decode(encoding)
    except UnicodeDecodeError:
        text = raw.decode("cp1252")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    csv.field_size_limit(1024 * 1024)
    reader = csv.reader(io.StringIO(text, newline=""), dialect)
    headers = next(reader, [])
    if len(headers) > 500:
        raise ValueError("CSV exceeds the 500 column processing limit.")
    records = []
    maximum_columns = len(headers)
    for row_num, row in enumerate(reader, 2):
        if row_num > MAX_ROWS:
            raise ValueError("CSV exceeds the 50,000 row processing limit.")
        if not any(row):
            continue
        if len(row) > 500:
            raise ValueError("CSV exceeds the 500 column processing limit.")
        maximum_columns = max(maximum_columns, len(row))
        row_headers = [headers[i] if i < len(headers) and headers[i] else get_column_letter(i + 1) for i in range(len(row))]
        context = " | ".join(f"{header}: {value}" for header, value in zip(row_headers, row))
        for column, value in enumerate(row):
            if value:
                records.append(cell_record("csv_cell", value, row_headers[column], row_num, get_column_letter(column + 1), f"{get_column_letter(column + 1)}{row_num}", None, {"header_row": 1, "physical_line": reader.line_num}, context))
                if len(records) > MAX_RECORDS:
                    raise ValueError("CSV exceeds the 100,000 searchable record limit.")
    if records:
        records[0]["metadata_json"]["table_headers"] = [{"key": get_column_letter(column + 1), "label": headers[column] if column < len(headers) and headers[column] else get_column_letter(column + 1)} for column in range(maximum_columns)]
    return records, [], []


def parse_pdf(path):
    records, pages = [], []
    with fitz.open(path) as document:
        if document.needs_pass:
            raise ValueError("Password-protected PDFs are not supported.")
        if len(document) > 2000:
            raise ValueError("PDF exceeds the 2,000 page processing limit.")
        for page_number, page in enumerate(document, 1):
            full_text = page.get_text("text")
            pages.append(dict(page_number=page_number, text=full_text, width=page.rect.width, height=page.rect.height))
            for block_number, block in enumerate(page.get_text("blocks"), 1):
                text = block[4].strip()
                if len(block) > 6 and block[6] != 0 or not text:
                    continue
                records.append(dict(record_type="pdf_block", content=text, value=text, page_number=page_number, metadata_json={"block": block_number, "bbox": [round(float(v), 2) for v in block[:4]]}, context=""))
            # Tables retain exact page and table coordinates in addition to text blocks.
            try:
                tables = page.find_tables().tables
            except Exception:
                tables = []
            for table_number, table in enumerate(tables, 1):
                rows = table.extract()
                for row_number, row in enumerate(rows, 1):
                    text = " | ".join(str(v or "") for v in row)
                    if text.strip(" |"):
                        records.append(dict(record_type="pdf_table", content=text, value=text, page_number=page_number, metadata_json={"table": table_number, "table_row": row_number, "bbox": list(table.bbox), "values": row}, context=""))
            if len(records) > MAX_RECORDS:
                raise ValueError("PDF exceeds the 100,000 searchable record limit.")
    if not records:
        raise ValueError("No searchable text was found. Image-only PDFs require a text layer; OCR is not enabled.")
    return records, pages, []


def parse_pptx(path):
    presentation = Presentation(path)
    records, slides = [], []
    if len(presentation.slides) > 2000:
        raise ValueError("Presentation exceeds the 2,000 slide processing limit.")
    for slide_number, slide in enumerate(presentation.slides, 1):
        title = slide.shapes.title.text if slide.shapes.title else f"Slide {slide_number}"
        parts = []
        for block_number, shape in enumerate(slide.shapes, 1):
            if shape.has_text_frame and shape.text.strip():
                text = shape.text.strip()
                parts.append(text)
                records.append(dict(record_type="slide_text", content=text, value=text, slide_number=slide_number, metadata_json={"block": block_number, "title": title, "shape": shape.name}, context=title))
            if shape.has_table:
                for row_number, row in enumerate(shape.table.rows, 1):
                    text = " | ".join(cell.text for cell in row.cells)
                    parts.append(text)
                    records.append(dict(record_type="slide_table", content=text, value=text, slide_number=slide_number, metadata_json={"block": block_number, "title": title, "table_row": row_number}, context=title))
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                records.append(dict(record_type="slide_notes", content=notes, value=notes, slide_number=slide_number, metadata_json={"title": title, "block": "notes"}, context=title))
        slides.append(dict(slide_number=slide_number, title=title, content="\n\n".join(parts), notes=notes))
        if len(records) > MAX_RECORDS:
            raise ValueError("Presentation exceeds the 100,000 searchable record limit.")
    return records, [], slides


def parse_file(path, file_type):
    validate_container(path, file_type)
    records, pages, slides = {"xlsx": parse_xlsx, "csv": parse_csv, "pdf": parse_pdf, "pptx": parse_pptx}[file_type](path)
    if not records:
        raise ValueError("The document contains no searchable records.")
    return records, pages, slides
