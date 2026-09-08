"""Exercise a running MCCIA deployment through its public HTTP API.

Creates real XLSX, CSV, PDF, and PPTX documents, then removes only the workspace
created by this run. Install backend/requirements.txt before running this script.
"""
from __future__ import annotations

import argparse
import csv
import io
import os
from pathlib import Path
import tempfile
import time
import uuid

import pymupdf as fitz
import httpx
from openpyxl import Workbook
from pptx import Presentation


def create_documents(directory: Path, marker: str) -> list[Path]:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "West Region"
    sheet.append(["Supplier", "Total Spending"])
    sheet.append([f"Approved supplier {marker}", 1845000])
    workbook.save(directory / f"{marker}_spending.xlsx")
    workbook.close()

    with (directory / f"{marker}_suppliers.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Supplier", "Purchase Order", "Approval Status"])
        writer.writerow([f"Industrial supplier {marker}", "PO-1023", "Approved"])

    with fitz.open() as document:
        document.new_page().insert_text((72, 72), "Compliance register: cover page")
        document.new_page().insert_text((72, 72), f"Factory licence renewal {marker}")
        document.save(directory / f"{marker}_compliance.pdf")

    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[0]).shapes.title.text = "Operations review"
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Machine maintenance"
    slide.placeholders[1].text = f"CNC inspection complete {marker}"
    presentation.save(directory / f"{marker}_operations.pptx")

    return sorted(directory.iterdir())


def checked(response: httpx.Response, status: int | tuple[int, ...] = (200, 201)) -> httpx.Response:
    statuses = (status,) if isinstance(status, int) else status
    if response.status_code not in statuses:
        raise AssertionError(f"{response.request.method} {response.request.url.path}: expected {statuses}, got {response.status_code}: {response.text[:800]}")
    return response


def run(base_url: str, email: str, password: str) -> None:
    marker = "smoke" + uuid.uuid4().hex[:12]
    workspace_id = None
    with httpx.Client(base_url=base_url.rstrip("/"), headers={"Origin": base_url.rstrip("/")}, timeout=120, follow_redirects=True, trust_env=False) as client:
        checked(client.get("/api/health", timeout=10))
        print("PASS service health")
        checked(client.get("/api/workspaces"), 401)
        print("PASS authentication required")
        login = checked(client.post("/api/auth/login", json={"email": email, "password": password})).json()
        user = login.get("user", login)
        if user.get("role") != "Admin":
            raise AssertionError("The smoke test needs an administrator to create and clean up its isolated workspace.")
        print("PASS login and session cookie")

        try:
            workspace = checked(client.post("/api/workspaces", json={"name": f"Verification {marker}", "description": "Temporary automated HTTP smoke test"})).json()
            workspace_id = workspace["id"]
            print("PASS workspace creation")
            checked(client.get("/api/auth/options"))
            checked(client.get("/api/auth/me"))
            checked(client.get("/api/settings"))
            checked(client.get("/api/users"))
            checked(client.get(f"/api/workspaces/{workspace_id}/members"))
            checked(client.get("/api/dashboard", params={"workspace_id": workspace_id}))
            checked(client.get("/api/search/history", params={"workspace_id": workspace_id}))
            print("PASS dashboard, settings, users, membership, and search history APIs")

            with tempfile.TemporaryDirectory(prefix="mccia-smoke-") as temporary:
                paths = create_documents(Path(temporary), marker)
                multipart = [("files", (path.name, path.read_bytes(), "application/octet-stream")) for path in paths]
                uploaded = checked(client.post("/api/files/upload", data={"workspace_id": str(workspace_id), "duplicate": "skip"}, files=multipart), 202).json()
                assert not uploaded.get("errors"), uploaded
                deadline = time.monotonic() + 120
                while True:
                    files = checked(client.get("/api/files", params={"workspace_id": workspace_id})).json()
                    if isinstance(files, dict):
                        files = files["files"]
                    failed = [item for item in files if item["processing_status"] == "failed"]
                    assert not failed, f"Parsing failed: {failed}"
                    if len(files) == 4 and all(item["processing_status"] == "indexed" for item in files):
                        break
                    if time.monotonic() >= deadline:
                        raise AssertionError(f"Indexing did not finish within 120 seconds: {files}")
                    time.sleep(0.5)
                print("PASS upload of four real document formats")

                query_terms = {"xlsx": "supplier", "csv": "industrial", "pdf": "licence", "pptx": "CNC"}
                for file_type in ("xlsx", "csv", "pdf", "pptx"):
                    body = {"query": f"{marker} {query_terms[file_type]}", "workspace_id": workspace_id, "filters": {"file_type": file_type}}
                    result = checked(client.post("/api/search", json=body)).json()
                    evidence = result["results"]
                    if not evidence:
                        raise AssertionError(f"No indexed evidence for uploaded {file_type}")
                    primary = evidence[0]
                    assert primary["file"]["type"] == file_type, primary
                    location = primary["location"]
                    if file_type == "xlsx":
                        assert location.get("sheet") == "West Region" and location.get("row") == 2 and location.get("cell"), location
                    elif file_type == "csv":
                        assert location.get("row") == 2 and location.get("column"), location
                    elif file_type == "pdf":
                        assert location.get("page") == 2, location
                    else:
                        assert location.get("slide") == 2, location
                    source = checked(client.get(f"/api/records/{primary['result_id']}/source"))
                    assert source.json(), "The source viewer returned an empty payload."
                    original = checked(client.get(f"/api/files/{primary['file']['id']}/download"))
                    assert original.content, "The original document download is empty."
                    print(f"PASS {file_type.upper()} search, precise citation, source, and original download")

                export = checked(client.post("/api/search/export", json={"query": marker, "workspace_id": workspace_id, "filters": {}}))
                rows = list(csv.reader(io.StringIO(export.content.decode("utf-8-sig"))))
                assert len(rows) > 4, "Expected header and evidence from four file formats in CSV export."
                assert marker in export.text, "Export is missing the matching evidence."
                print("PASS CSV evidence export")

                audit = checked(client.get("/api/audit", params={"workspace_id": workspace_id}))
                assert marker in audit.text, "Audit history is missing the test upload/search context."
                print("PASS persisted audit history")
        finally:
            if workspace_id is not None:
                checked(client.delete(f"/api/workspaces/{workspace_id}"), (200, 204))
                print("PASS temporary workspace cleanup")
        checked(client.post("/api/auth/logout"), (200, 204))
        checked(client.get("/api/workspaces"), 401)
        print("PASS logout invalidates session")
    print(f"All HTTP smoke checks passed against {base_url}.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--url", default="http://localhost:3000", help="Frontend URL (default) or backend URL; do not append /api.")
    parser.add_argument("--email", default=os.getenv("ADMIN_EMAIL", "admin@mccia.org"))
    parser.add_argument("--password", default=os.getenv("ADMIN_PASSWORD", "Mccia@2026!"))
    arguments = parser.parse_args()
    run(arguments.url, arguments.email, arguments.password)
